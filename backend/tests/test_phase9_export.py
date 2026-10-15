import os
os.environ['MONGODB_DB_NAME'] = 'cognitrace_test_db'
import pytest
from app.services.export_service import ExportService
import csv
import json
import io
from openpyxl import load_workbook

def test_export_service_csv():
    data = [
        {"StudentID": "S1", "Concept": "Arrays", "Score": 10},
        {"StudentID": "S2", "Concept": "Loops", "Score": None}
    ]
    fieldnames = ["StudentID", "Concept", "Score"]
    
    csv_bytes = ExportService.generate_csv(data, fieldnames)
    assert isinstance(csv_bytes, bytes)
    
    reader = csv.DictReader(io.StringIO(csv_bytes.decode('utf-8')))
    rows = list(reader)
    assert len(rows) == 2
    assert reader.fieldnames == fieldnames
    
    assert rows[0]["StudentID"] == "S1"
    assert rows[1]["Score"] == ""

def test_export_service_json():
    data = [
        {"StudentID": "S1", "Concept": "Arrays", "InternalID": 99},
        {"StudentID": "S2", "Concept": "Loops", "InternalID": 88}
    ]
    fieldnames = ["StudentID", "Concept"]
    
    json_bytes = ExportService.generate_json(data, fieldnames)
    parsed = json.loads(json_bytes.decode('utf-8'))
    
    assert len(parsed) == 2
    assert "InternalID" not in parsed[0]
    assert parsed[0]["StudentID"] == "S1"
    assert parsed[0]["Concept"] == "Arrays"

def test_export_service_excel():
    data = [
        {"StudentID": "S1", "Concept": "=cmd|' /C calc'!A0", "Score": 10}
    ]
    fieldnames = ["StudentID", "Concept", "Score"]
    
    excel_bytes = ExportService.generate_excel(data, fieldnames)
    assert isinstance(excel_bytes, bytes)
    
    wb = load_workbook(filename=io.BytesIO(excel_bytes))
    ws = wb.active
    assert ws.title == "Dataset"
    
    headers = [cell.value for cell in ws[1]]
    assert headers == fieldnames
    
    # Check formula escape
    concept_val = ws.cell(row=2, column=2).value
    assert concept_val.startswith("'=")
